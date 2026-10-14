"""
DataMind V4.1
统一文件加载与数据归一化

职责：

1. 按扩展名识别上传文件的类型
2. 把各种格式统一读成 pandas DataFrame
     - 表格类：csv / txt / tsv / xlsx / xlsm / xls
     - Word  ：docx（抽取其中的表格）
     - 图片  ：png / jpg / jpeg / bmp / webp / gif
               交给视觉模型把图片里的表格或图表转成数据
3. 归一化：把读出来的数据统一写成一份标准 CSV

为什么要归一化：

    分析链路（Planner / Executor / run_stat_calc / Chart Executor）
    只认 CSV 路径。

    只要在入口处把任何格式都变成标准 CSV，
    下游一行都不用改，Excel / Word / 图片 就自动被支持了。

4. 遇到无法变成表格的输入（纯文字段落、风景照等），
   返回 text 类型，由报告节点直接根据内容回答用户问题。
"""

import base64
import io
import os
import re
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from services.llm_service import call_llm


# ========================================
# 支持的文件类型
# ========================================

TABLE_EXTENSIONS = {
    ".csv": "CSV 文件",
    ".txt": "文本表格",
    ".tsv": "TSV 文件",
    ".xlsx": "Excel 文件",
    ".xlsm": "Excel 文件",
    ".xls": "Excel 文件（旧版）",
}

WORD_EXTENSIONS = {
    ".docx": "Word 文档",
}

IMAGE_EXTENSIONS = {
    ".png": "图片",
    ".jpg": "图片",
    ".jpeg": "图片",
    ".bmp": "图片",
    ".webp": "图片",
    ".gif": "图片",
}

SUPPORTED_EXTENSIONS = {}
SUPPORTED_EXTENSIONS.update(TABLE_EXTENSIONS)
SUPPORTED_EXTENSIONS.update(WORD_EXTENSIONS)
SUPPORTED_EXTENSIONS.update(IMAGE_EXTENSIONS)


# ========================================
# 限制
# ========================================

# 归一化 CSV 的存放目录
NORMALIZED_DIR = "temp"

# 交给视觉模型前，图片最长边压缩到这个尺寸（控制 token 消耗）
MAX_IMAGE_SIDE = 1600

# 图片大小上限
MAX_IMAGE_BYTES = 8 * 1024 * 1024

# 纯文本内容进入 prompt 的长度上限
MAX_CONTEXT_CHARS = 8000


# ========================================
# 图片提取提示词
#
# 分两种情况，已经被实测验证：
#   - 表格截图 -> 完整 CSV（连 dtype 都对）
#   - 柱状图   -> 柱子上的数值被正确读出
#   - 纯文字   -> 返回 NO_TABLE + 描述
# ========================================

IMAGE_EXTRACT_PROMPT = """你是一个数据提取助手。请仔细观察这张图片，然后按下面的规则输出。

情况 A：图片中包含表格，或包含带数值的图表（柱状图、折线图、饼图等）。
请把其中的数据提取成 CSV：
- 第一行是表头，之后每行是一条记录
- 只输出 CSV 本身，不要 Markdown 代码块，不要任何解释
- 数值不要带千分位逗号、货币符号或单位
- 列名保持图片里的原始文字

情况 B：图片中没有可以结构化的数据（例如风景照、纯文字段落）。
第一行只输出：NO_TABLE
从第二行开始，用中文简要描述图片内容，不超过 200 字。
"""


# ========================================
# 基础工具
# ========================================

def get_extension(path: str) -> str:
    """取小写扩展名。"""

    return os.path.splitext(path or "")[1].lower()


def source_label(extension: str) -> str:
    """扩展名 -> 中文类型名。"""

    return SUPPORTED_EXTENSIONS.get(extension, "未知格式")


def is_supported(path: str) -> bool:
    return get_extension(path) in SUPPORTED_EXTENSIONS


def supported_hint() -> str:
    """给用户看的支持格式说明。"""

    return "、".join(sorted(SUPPORTED_EXTENSIONS))


def accepted_extensions() -> List[str]:
    """给前端 input accept 用的扩展名列表。"""

    return sorted(SUPPORTED_EXTENSIONS)


# ========================================
# 表格类
# ========================================

def read_csv_robust(path: str, **kwargs) -> pd.DataFrame:
    """
    依次尝试常见编码读取 CSV。

    国内用户导出的 CSV 经常是 GBK，
    直接按 UTF-8 读会直接抛 UnicodeDecodeError。
    """

    last_error = None

    for encoding in ("utf-8", "utf-8-sig", "gbk", "latin-1"):

        try:

            return pd.read_csv(
                path,
                encoding=encoding,
                **kwargs
            )

        except UnicodeDecodeError as exc:

            last_error = exc

    if last_error is not None:

        raise last_error

    raise ValueError("无法读取该 CSV 文件")


def read_excel_best_sheet(
    path: str,
) -> Tuple[pd.DataFrame, List[str]]:
    """
    读取 Excel，并挑选行数最多的工作表。

    多工作表时会给出一条提示，
    避免用户以为「明明有数据却读不到」。
    """

    sheets = pd.read_excel(
        path,
        sheet_name=None
    )

    notes: List[str] = []

    if not sheets:

        raise ValueError("Excel 文件里没有任何工作表")

    names = list(sheets.keys())

    best_name = max(
        names,
        key=lambda name: len(sheets[name])
    )

    if len(names) > 1:

        notes.append(
            f"Excel 共有 {len(names)} 个工作表，"
            f"已使用行数最多的「{best_name}」"
        )

    return sheets[best_name], notes


def clean_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """
    去掉整行/整列的空值，并规整列名。
    """

    df = df.dropna(how="all")
    df = df.dropna(axis=1, how="all")

    df.columns = [
        str(column).strip()
        for column in df.columns
    ]

    return df.reset_index(drop=True)


def dedupe_headers(header: List[str]) -> List[str]:
    """
    Word 表格里经常出现空表头或重复表头，
    这里补成唯一列名，否则 DataFrame 会出现重名列。
    """

    result: List[str] = []
    seen: Dict[str, int] = {}

    for index, name in enumerate(header):

        name = (name or "").strip() or f"列{index + 1}"

        if name in seen:

            seen[name] += 1
            name = f"{name}_{seen[name]}"

        else:

            seen[name] = 1

        result.append(name)

    return result


def load_table(path: str) -> Tuple[pd.DataFrame, List[str]]:
    """
    读取表格类文件（CSV / TXT / TSV / Excel）。
    """

    extension = get_extension(path)
    notes: List[str] = []

    if extension == ".csv":

        df = read_csv_robust(path)

    elif extension == ".tsv":

        df = read_csv_robust(path, sep="\t")

    elif extension == ".txt":

        # 文本表格：让 pandas 自己猜分隔符
        df = read_csv_robust(
            path,
            sep=None,
            engine="python"
        )

    elif extension in (".xlsx", ".xlsm", ".xls"):

        df, notes = read_excel_best_sheet(path)

    else:

        raise ValueError(f"不支持的表格格式：{extension}")

    return clean_dataframe(df), notes


# ========================================
# Word 文档
# ========================================

def load_word(
    path: str,
) -> Tuple[Optional[pd.DataFrame], str, List[str]]:
    """
    读取 Word 文档。

    返回：

        (表格 DataFrame 或 None, 正文文本, 提示)
    """

    from docx import Document

    document = Document(path)

    notes: List[str] = []

    tables: List[pd.DataFrame] = []

    for table in document.tables:

        rows = [
            [cell.text.strip() for cell in row.cells]
            for row in table.rows
        ]

        rows = [
            row for row in rows
            if any(cell for cell in row)
        ]

        if len(rows) < 2:

            continue

        header = dedupe_headers(rows[0])

        # 合并单元格可能让某一行少几个格子，
        # 这里统一补齐/截断，避免 pandas 直接报错。
        width = len(header)

        body = []

        for row in rows[1:]:

            if len(row) < width:

                row = row + [""] * (width - len(row))

            body.append(row[:width])

        tables.append(
            pd.DataFrame(body, columns=header)
        )

    text = "\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )

    if not tables:

        notes.append("文档里没有找到可用的表格")

        return None, text, notes

    best = max(tables, key=len)

    if len(tables) > 1:

        notes.append(
            f"文档里共有 {len(tables)} 个表格，"
            f"已使用行数最多的那个（{len(best)} 行）"
        )

    return clean_dataframe(best), text, notes


# ========================================
# 图片
# ========================================

def encode_image(path: str) -> str:
    """
    把图片压到合适尺寸后转成 data URL。

    压缩是为了控制视觉模型的 token 消耗 ——
    手机截图动辄 4000 像素宽，直接发又贵又容易超限。
    """

    from PIL import Image

    image = Image.open(path)
    image = image.convert("RGB")

    longest = max(image.size)

    if longest > MAX_IMAGE_SIDE:

        ratio = MAX_IMAGE_SIDE / longest

        image = image.resize(
            (
                max(1, int(image.width * ratio)),
                max(1, int(image.height * ratio)),
            )
        )

    buffer = io.BytesIO()
    image.save(buffer, format="PNG")

    encoded = base64.b64encode(
        buffer.getvalue()
    ).decode()

    return "data:image/png;base64," + encoded


def parse_image_response(
    text: str,
) -> Tuple[Optional[pd.DataFrame], Optional[str], List[str]]:
    """
    解析视觉模型的输出。

    可能是 CSV，也可能是 NO_TABLE + 描述。
    """

    cleaned = (text or "").strip()

    if cleaned.startswith("```"):

        cleaned = re.sub(
            r"^```[a-zA-Z]*\s*",
            "",
            cleaned
        )

        cleaned = re.sub(r"\s*```$", "", cleaned).strip()

    if cleaned.upper().startswith("NO_TABLE"):

        description = cleaned[len("NO_TABLE"):].strip()

        return (
            None,
            description,
            ["图片里没有可结构化的数据，已转为内容解读"],
        )

    try:

        df = pd.read_csv(io.StringIO(cleaned))

    except Exception as exc:  # noqa: BLE001

        return (
            None,
            cleaned,
            [
                "图片识别结果无法解析成表格"
                f"（{type(exc).__name__}），已转为内容解读"
            ],
        )

    if df.empty:

        return (
            None,
            cleaned,
            ["图片识别结果为空，已转为内容解读"],
        )

    return (
        clean_dataframe(df),
        None,
        ["数据来自图片识别，可能存在误差，请核对关键数值"],
    )


def load_image(
    path: str,
) -> Tuple[Optional[pd.DataFrame], Optional[str], List[str]]:
    """
    用视觉模型把图片转成数据。
    """

    size = os.path.getsize(path)

    if size > MAX_IMAGE_BYTES:

        raise ValueError(
            f"图片太大（{size / 1024 / 1024:.1f} MB），"
            f"请压缩到 {MAX_IMAGE_BYTES // 1024 // 1024} MB 以内"
        )

    data_url = encode_image(path)

    messages = [
        {
            "role": "user",
            "content": [
                {
                    "type": "text",
                    "text": IMAGE_EXTRACT_PROMPT
                },
                {
                    "type": "image_url",
                    "image_url": {"url": data_url}
                },
            ],
        }
    ]

    response = call_llm(messages, temperature=0.0)

    return parse_image_response(response)


# ========================================
# 归一化输出
# ========================================

def write_normalized_csv(
    df: pd.DataFrame,
    output_dir: str = NORMALIZED_DIR,
) -> str:
    """
    把 DataFrame 写成标准 CSV，返回路径。

    分析链路只认这个文件。
    """

    os.makedirs(output_dir, exist_ok=True)

    filename = (
        "data_"
        + datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        + ".csv"
    )

    path = os.path.join(output_dir, filename)

    df.to_csv(
        path,
        index=False,
        encoding="utf-8"
    )

    return path


def clip_context(text: Optional[str]) -> Optional[str]:
    """限制纯文本内容的长度，避免 prompt 过长。"""

    if not text:

        return text

    if len(text) <= MAX_CONTEXT_CHARS:

        return text

    return (
        text[:MAX_CONTEXT_CHARS]
        + "\n…（内容过长，已截断）"
    )


# ========================================
# 主入口
# ========================================

def ingest(
    path: str,
    output_dir: str = NORMALIZED_DIR,
) -> Dict[str, Any]:
    """
    读取任意支持的文件，返回统一结构。

    返回字段：

        kind            table / text / error
        source_type     csv / excel / word_table / word_text /
                        image_table / image_text
        normalized_path 归一化后的 CSV 路径（仅 table）
        shape           [行数, 列数]（仅 table）
        dataset_info    给下游提示用的数据来源说明
        context         非表格内容的正文（仅 text）
        notes           给用户看的提示
        error           失败原因（仅 error）
    """

    extension = get_extension(path)
    label = source_label(extension)

    result: Dict[str, Any] = {
        "kind": "error",
        "source_type": None,
        "normalized_path": None,
        "shape": None,
        "dataset_info": None,
        "context": None,
        "notes": [],
        "error": None,
    }

    if not is_supported(path):

        result["error"] = (
            f"暂不支持的文件格式：{extension or '（无扩展名）'}。"
            f"支持的格式：{supported_hint()}"
        )

        return result

    notes: List[str] = []

    try:

        if extension in IMAGE_EXTENSIONS:

            df, context, extra = load_image(path)

            source_type = (
                "image_table"
                if df is not None
                else "image_text"
            )

        elif extension in WORD_EXTENSIONS:

            df, context, extra = load_word(path)

            source_type = (
                "word_table"
                if df is not None
                else "word_text"
            )

        else:

            df, extra = load_table(path)

            context = None

            source_type = (
                "excel"
                if extension in (".xlsx", ".xlsm", ".xls")
                else "csv"
            )

        notes.extend(extra)

    except Exception as exc:  # noqa: BLE001

        result["source_type"] = source_type_of(extension)
        result["error"] = (
            f"读取文件失败：{type(exc).__name__}: {exc}"
        )

        return result

    result["source_type"] = source_type
    result["notes"] = notes

    # ======================================
    # 没有表格 -> 内容解读
    # ======================================

    if df is None:

        result["kind"] = "text"
        result["context"] = clip_context(context or "")

        return result

    if df.empty:

        result["error"] = "文件里没有读到任何数据"

        return result

    # ======================================
    # 有表格 -> 归一化成标准 CSV
    # ======================================

    normalized = write_normalized_csv(df, output_dir)

    result["kind"] = "table"
    result["normalized_path"] = normalized
    result["shape"] = [int(df.shape[0]), int(df.shape[1])]
    result["dataset_info"] = (
        f"=== 数据来源 ===\n{label}\n"
    )

    return result


def source_type_of(extension: str) -> Optional[str]:
    """异常情况下也要能报告来源类型。"""

    if extension in IMAGE_EXTENSIONS:
        return "image"

    if extension in WORD_EXTENSIONS:
        return "word"

    if extension in TABLE_EXTENSIONS:
        return "excel" if extension in (".xlsx", ".xlsm", ".xls") else "csv"

    return None
