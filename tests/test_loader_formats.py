"""
多格式输入（V4.1）：CSV / Excel / Word / 图片 -> 统一归一化成标准 CSV

由 ``_scratch/v41_verify.py`` 移植而来，覆盖三块：

1. ``tools/loader.py`` 的逐格式解析与归一化
   （``ingest`` / ``read_csv_robust`` / ``read_excel_best_sheet`` /
   ``load_word`` / ``load_image`` / ``parse_image_response`` /
   ``write_normalized_csv``）

2. ``agent/nodes/dataset.py`` 的路由
   表格 -> 分析链路（retrieve -> planner -> ...）
   Word 正文 / 图片解读 -> 直接进报告节点
   读取失败 -> 进报告节点如实说明

3. 完整图（``agent.graph``）与 HTTP 契约（``POST /analyze``）的对应行为

与原脚本的差异：

1. 用 ``assert`` 取代 ``print`` 报告 + ``sys.exit``，不再有自定义 ``check()`` 收集器；
2. 图片走的是视觉模型，这里**全部用 ``stub_llm`` 伪造**：
   源脚本里那两次真实的视觉调用被换成桩，测试完全离线；
3. 所有中间产物（归一化 CSV、上传文件）都落在 ``tmp_path`` 里，
   不往项目目录写任何东西；
4. 不依赖当前工作目录，项目根目录由本文件路径推导（与 ``tests/conftest.py`` 一致）。

本文件不发出任何真实模型请求：conftest 里的 ``no_real_llm`` 会拦截漏网调用。
"""

import json
import os
import sys

import pandas as pd
import pytest

# =========================================================
# 项目根目录
# =========================================================

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT = os.path.dirname(HERE)
if not os.path.isdir(os.path.join(PROJECT, "agent")):
    PROJECT = os.path.dirname(PROJECT)

if PROJECT not in sys.path:
    sys.path.insert(0, PROJECT)

from agent.graph import build_graph, route_after_dataset  # noqa: E402
from agent.nodes.dataset import analyze_dataset  # noqa: E402
from tools.loader import (  # noqa: E402
    NORMALIZED_DIR,
    dedupe_headers,
    ingest,
    load_image,
    load_word,
    parse_image_response,
    read_csv_robust,
    read_excel_best_sheet,
    supported_hint,
    write_normalized_csv,
)


# =========================================================
# 公共数据
# =========================================================

SALES_COLUMNS = ["地区", "季度", "销售额", "利润"]

SALES_ROWS = [
    ["华东", "2026Q1", 125000, 24000],
    ["华南", "2026Q1", 98000, 17500],
    ["华北", "2026Q1", 76500, 13000],
    ["华东", "2026Q2", 143200, 28900],
]

# 4 行销售额合计；用来证明下游工具真的读了归一化后的那份 CSV
SALES_TOTAL = "442,700"

# Word 正文里那句可辨识的话
PROSE_LINE = "决定把数据迁移推迟到明年三月。"

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_DOCX = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

# 视觉模型的桩回复
IMAGE_TABLE_REPLY = (
    "地区,季度,销售额\n"
    "华东,2026Q1,125000\n"
    "华南,2026Q1,98000\n"
    "华北,2026Q1,76500\n"
)

IMAGE_PROSE_REPLY = (
    "NO_TABLE\n"
    "图片是一段会议纪要：决定把数据迁移推迟到明年三月，没有涉及数值。"
)


# =========================================================
# 夹具：不往项目目录写文件
# =========================================================

@pytest.fixture(autouse=True)
def loader_output_dir(monkeypatch, tmp_path):
    """
    把归一化 CSV 的输出目录重定向到 ``tmp_path``。

    ``ingest()`` 的 ``output_dir`` 默认值是模块常量 ``"temp"``（相对当前工作
    目录），不拦的话跑一次测试就会往项目里写 ``temp/data_*.csv``。
    显式传了别的目录的调用（例如 ``write_normalized_csv`` 自身的用例）
    照原样转发。
    """

    import tools.loader as loader

    target = tmp_path / "normalized"
    original = loader.write_normalized_csv

    def redirected(df, output_dir=NORMALIZED_DIR):

        if output_dir == NORMALIZED_DIR:

            output_dir = str(target)

        return original(df, output_dir)

    monkeypatch.setattr(loader, "write_normalized_csv", redirected)

    return str(target)


# =========================================================
# 夹具：各格式的样例文件
# =========================================================

@pytest.fixture
def sales_frame():
    return pd.DataFrame(SALES_ROWS, columns=SALES_COLUMNS)


@pytest.fixture
def csv_file(tmp_path, sales_frame):
    path = tmp_path / "sales.csv"
    sales_frame.to_csv(path, index=False, encoding="utf-8")
    return str(path)


@pytest.fixture
def gbk_csv_file(tmp_path, sales_frame):
    path = tmp_path / "sales_gbk.csv"
    sales_frame.to_csv(path, index=False, encoding="gbk")
    return str(path)


@pytest.fixture
def xlsx_file(tmp_path):
    """
    两个工作表：小的「说明」+ 大的「明细」。

    用来验证「挑行数最多的工作表」以及多表提示。
    """

    from openpyxl import Workbook

    workbook = Workbook()

    summary = workbook.active
    summary.title = "说明"
    summary.append(SALES_COLUMNS)
    for row in SALES_ROWS[:2]:
        summary.append(row)

    detail = workbook.create_sheet("明细")
    detail.append(SALES_COLUMNS)
    for row in SALES_ROWS:
        detail.append(row)

    path = tmp_path / "sales.xlsx"
    workbook.save(path)

    return str(path)


@pytest.fixture
def docx_table_file(tmp_path):
    """带真表格 + 段落的 Word 文档。"""

    from docx import Document

    document = Document()
    document.add_heading("季度销售汇总", level=1)
    document.add_paragraph("以下是本季度各地区的销售情况。")

    table = document.add_table(rows=1, cols=3)
    for index, title in enumerate(["地区", "销售额", "利润"]):
        table.rows[0].cells[index].text = title

    for row in SALES_ROWS[:3]:
        cells = table.add_row().cells
        cells[0].text = row[0]
        cells[1].text = str(row[2])
        cells[2].text = str(row[3])

    path = tmp_path / "report_with_table.docx"
    document.save(path)

    return str(path)


@pytest.fixture
def docx_prose_file(tmp_path):
    """只有正文、没有表格的 Word 文档。"""

    from docx import Document

    document = Document()
    document.add_heading("会议纪要", level=1)
    document.add_paragraph("讨论了下一季度的排期。")
    document.add_paragraph(PROSE_LINE)
    document.add_paragraph("本次会议没有涉及任何数值。")

    path = tmp_path / "prose.docx"
    document.save(path)

    return str(path)


@pytest.fixture
def pdf_file(tmp_path):
    """不支持的格式：扩展名对就行，内容无所谓。"""

    path = tmp_path / "notes.pdf"
    path.write_bytes(b"%PDF-1.4 not a real pdf")

    return str(path)


def _make_png(path, lines):
    """
    造一张合法 PNG。

    视觉调用被 ``stub_llm`` 换成了桩，所以图片内容不会被任何模型读到；
    这里只需要保证 PIL 打开 / 压缩 / base64 这条链路能真正跑通。
    """

    from PIL import Image, ImageDraw

    image = Image.new("RGB", (620, 200), "white")
    draw = ImageDraw.Draw(image)

    for index, line in enumerate(lines):
        draw.text((20, 20 + index * 30), line, fill="black")

    image.save(path)

    return path


@pytest.fixture
def image_table_file(tmp_path):
    """「截图里是一张表」的图片。"""

    return str(_make_png(tmp_path / "table.png", [
        "Region   Quarter   Sales",
        "East     2026Q1    125000",
        "South    2026Q1    98000",
        "North    2026Q1    76500",
    ]))


@pytest.fixture
def image_prose_file(tmp_path):
    """「截图里只有文字」的图片。"""

    return str(_make_png(tmp_path / "prose.png", [
        "Meeting notes",
        "We agreed to postpone the migration.",
        "No figures were reviewed.",
    ]))


# =========================================================
# B. loader：每种格式
# =========================================================

TABLE_INPUTS = [
    ("csv_file", "csv", [4, 4]),
    ("gbk_csv_file", "csv", [4, 4]),
    ("xlsx_file", "excel", [4, 4]),
    ("docx_table_file", "word_table", [3, 3]),
]


@pytest.mark.parametrize("fixture_name, source_type, shape", TABLE_INPUTS)
def test_table_inputs_normalize_to_a_standard_csv(
    fixture_name, source_type, shape, request, tmp_path
):
    path = request.getfixturevalue(fixture_name)

    result = ingest(path)

    assert result["kind"] == "table", result.get("error")
    assert result["source_type"] == source_type
    assert result["shape"] == shape, result["shape"]

    normalized = result["normalized_path"]

    assert normalized, "table 结果必须带 normalized_path"
    assert os.path.isfile(normalized), normalized

    # 归一化文件落在 tmp_path 里 —— 项目目录没有被写过（见 loader_output_dir）
    assert os.path.abspath(normalized).startswith(str(tmp_path))

    # 归一化后的 CSV 重新读回来还是同样的形状
    assert list(pd.read_csv(normalized).shape) == shape


def test_read_csv_robust_falls_back_to_gbk(gbk_csv_file, sales_frame):
    frame = read_csv_robust(gbk_csv_file)

    assert list(frame.columns) == SALES_COLUMNS
    assert len(frame) == len(sales_frame)


def test_gbk_csv_content_survives_normalization(gbk_csv_file, sales_frame):
    result = ingest(gbk_csv_file)

    assert result["kind"] == "table", result.get("error")

    back = pd.read_csv(result["normalized_path"])

    assert list(back.columns) == list(sales_frame.columns)
    assert back["销售额"].tolist() == sales_frame["销售额"].tolist()


def test_excel_ingest_uses_the_sheet_with_the_most_rows(xlsx_file):
    result = ingest(xlsx_file)

    assert result["kind"] == "table", result.get("error")
    assert result["source_type"] == "excel"
    assert result["shape"] == [4, 4], result["shape"]

    assert any("工作表" in note for note in result["notes"]), result["notes"]

    # 归一化出来的是行数多的「明细」，不是行数少的「说明」
    back = pd.read_csv(result["normalized_path"])

    assert back["季度"].tolist() == ["2026Q1", "2026Q1", "2026Q1", "2026Q2"]


def test_read_excel_best_sheet_names_the_sheet_it_picked(xlsx_file):
    frame, notes = read_excel_best_sheet(xlsx_file)

    assert len(frame) == 4
    assert notes, "多工作表必须给出提示"

    joined = " ".join(notes)

    assert "工作表" in joined
    assert "明细" in joined


def test_word_table_headers_and_values_survive(docx_table_file):
    result = ingest(docx_table_file)

    assert result["kind"] == "table", result.get("error")
    assert result["source_type"] == "word_table"
    assert result["shape"] == [3, 3], result["shape"]

    back = pd.read_csv(result["normalized_path"])

    assert list(back.columns) == ["地区", "销售额", "利润"]
    assert back.iloc[0].tolist() == ["华东", 125000, 24000]


def test_word_prose_becomes_text_and_says_why(docx_prose_file):
    result = ingest(docx_prose_file)

    assert result["kind"] == "text", result.get("error")
    assert result["source_type"] == "word_text"
    assert PROSE_LINE in (result["context"] or "")

    assert any(
        "没有找到可用的表格" in note for note in result["notes"]
    ), result["notes"]

    # 非表格内容不产出归一化 CSV
    assert result["normalized_path"] is None


def test_load_word_returns_the_prose_when_there_is_no_table(docx_prose_file):
    frame, text, notes = load_word(docx_prose_file)

    assert frame is None
    assert PROSE_LINE in text
    assert notes == ["文档里没有找到可用的表格"]


def test_unsupported_extension_reports_the_format_and_the_supported_list(pdf_file):
    result = ingest(pdf_file)

    assert result["kind"] == "error"
    assert result["source_type"] is None
    assert "pdf" in result["error"], result["error"]
    assert "csv" in result["error"], result["error"]
    assert supported_hint() in result["error"]


# ---------------------------------------------------------
# B7 / B8：图片（视觉模型被 stub_llm 伪造）
# ---------------------------------------------------------

def test_image_with_a_table_is_normalized_through_the_vision_model(
    image_table_file, stub_llm
):
    stub_llm.default = IMAGE_TABLE_REPLY

    result = ingest(image_table_file)

    assert result["kind"] == "table", result.get("error")
    assert result["source_type"] == "image_table"
    assert result["shape"] == [3, 3], result["shape"]

    back = pd.read_csv(result["normalized_path"])

    assert len(back) == 3
    assert "销售额" in " ".join(str(column) for column in back.columns)
    assert pd.api.types.is_numeric_dtype(back.iloc[:, -1]), back.dtypes.tolist()

    assert any(
        "图片识别" in note for note in result["notes"]
    ), result["notes"]

    # 图片确实是以 data URL 交给视觉模型的（模型本身是桩，不会发请求）
    assert len(stub_llm.calls) == 1, stub_llm.calls

    payload = stub_llm.prompts[0][-1]["content"]

    assert isinstance(payload, list)

    image_parts = [
        part for part in payload if part.get("type") == "image_url"
    ]

    assert len(image_parts) == 1
    assert image_parts[0]["image_url"]["url"].startswith("data:image/png;base64,")


def test_image_without_a_table_becomes_a_text_interpretation(
    image_prose_file, stub_llm
):
    stub_llm.default = IMAGE_PROSE_REPLY

    result = ingest(image_prose_file)

    assert result["kind"] == "text", result.get("error")
    assert result["source_type"] == "image_text"
    assert (result["context"] or "").strip()
    assert "推迟到明年三月" in result["context"]

    assert any(
        "没有可结构化的数据" in note for note in result["notes"]
    ), result["notes"]


def test_load_image_rejects_a_file_over_the_size_limit(image_table_file, monkeypatch):
    import tools.loader as loader

    monkeypatch.setattr(loader, "MAX_IMAGE_BYTES", 10)

    with pytest.raises(ValueError) as excinfo:
        load_image(image_table_file)

    assert "图片太大" in str(excinfo.value)


# 视觉模型的输出五花八门，parse_image_response 有四个出口：
# 解析成表格 / NO_TABLE 描述 / 解析失败 / 解析出来是空的
IMAGE_RESPONSE_CASES = [
    (
        "纯 CSV",
        "地区,销售额\n华东,125000\n华南,98000\n",
        "table",
    ),
    (
        "Markdown 代码块包着的 CSV",
        "```csv\n地区,销售额\n华东,125000\n```",
        "table",
    ),
    (
        "数据行缺字段（全空列会被清掉）",
        "地区,销售额,季度\n华东,125000\n",
        "table",
    ),
    (
        "NO_TABLE + 描述",
        "NO_TABLE\n这是一张风景照，没有数据。",
        "description",
    ),
    (
        "引号没闭合，读不出表",
        '地区,销售额\n"华东,125000\n',
        "unparsable",
    ),
    (
        "只有表头、没有数据行",
        "地区,销售额\n",
        "empty",
    ),
    (
        "一行纯文本（会被当成表头，所以是空的）",
        "没有表格的一段文字",
        "empty",
    ),
]


@pytest.mark.parametrize("label, text, expected", IMAGE_RESPONSE_CASES)
def test_parse_image_response_handles_every_shape(label, text, expected):
    frame, context, notes = parse_image_response(text)

    assert notes, label

    if expected == "table":

        assert frame is not None, label
        assert context is None, label
        assert any("图片识别" in note for note in notes), notes

    elif expected == "empty":

        assert frame is None, label
        assert any("为空" in note for note in notes), notes

    elif expected == "unparsable":

        assert frame is None, label
        assert any("无法解析成表格" in note for note in notes), notes

    else:

        assert frame is None, label
        assert context, label
        assert any("已转为内容解读" in note for note in notes), notes


def test_write_normalized_csv_writes_the_file_where_it_is_told(sales_frame, tmp_path):
    target = tmp_path / "out" / "normalized"

    path = write_normalized_csv(sales_frame, str(target))

    assert os.path.isfile(path)
    assert os.path.dirname(os.path.abspath(path)) == os.path.abspath(str(target))
    assert os.path.basename(path).startswith("data_")
    assert path.endswith(".csv")
    assert pd.read_csv(path).shape == sales_frame.shape


def test_dedupe_headers_fills_blanks_and_renames_duplicates():
    assert dedupe_headers(["", "地区", "地区", " 销售额 "]) == [
        "列1", "地区", "地区_2", "销售额",
    ]


# =========================================================
# C. 路由：dataset 节点 / 完整图
# =========================================================

def test_dataset_routes_a_spreadsheet_into_the_analysis_chain(
    xlsx_file, make_state, trace_events, trace_titles
):
    result = analyze_dataset(make_state(
        question="销售额是多少",
        file_path=xlsx_file,
        dataset_info=None,
    ))

    assert result["input_kind"] == "table"
    assert result["source_type"] == "excel"
    # V4.11：表格数据先进入数据体检，再检索口径
    assert route_after_dataset(result) == "health_check"

    # file_path 被改写指向归一化后的 CSV —— 下游只认 CSV
    assert result["file_path"].endswith(".csv")
    assert os.path.isfile(result["file_path"])

    assert "Excel 文件" in result["dataset_info"]
    assert "dataset_analyzed" in trace_events(result)
    assert any("Excel" in str(title) for title in trace_titles(result))


NON_TABLE_INPUTS = [
    ("docx_prose_file", "word_text"),
    ("image_prose_file", "image_text"),
]


@pytest.mark.parametrize("fixture_name, source_type", NON_TABLE_INPUTS)
def test_dataset_routes_non_table_content_straight_to_the_report(
    fixture_name, source_type, request, stub_llm, make_state, trace_events
):
    if fixture_name == "image_prose_file":

        stub_llm.default = IMAGE_PROSE_REPLY

    path = request.getfixturevalue(fixture_name)

    result = analyze_dataset(make_state(
        question="会议讲了什么",
        file_path=path,
        dataset_info=None,
    ))

    assert result["input_kind"] == "text"
    assert result["source_type"] == source_type
    assert result["dataset_info"] is None
    assert (result["input_context"] or "").strip()
    assert "content_extracted" in trace_events(result)
    assert route_after_dataset(result) == "report"


def test_dataset_routes_a_read_failure_to_the_report(
    pdf_file, make_state, trace_events
):
    result = analyze_dataset(make_state(
        question="销售额是多少",
        file_path=pdf_file,
        dataset_info=None,
    ))

    assert result["input_kind"] == "error"
    assert result["dataset_info"] is None
    assert result["error"]
    assert route_after_dataset(result) == "report"
    assert "dataset_error" in trace_events(result)


@pytest.fixture
def stubbed_graph(stub_llm, tool_model, fake_tool_call):
    """
    把整张图的模型侧全部换成桩：

    Planner / Reviewer / Chart Planner / Reporter 走 ``stub_llm``，
    Executor 的原生 Tool Calling 走 ``tool_model``。
    """

    stub_llm.reply("中的 Planner", json.dumps(
        {"plan": [{"step": 1, "task": "汇总销售额", "tool": "run_stat_calc"}]},
        ensure_ascii=False,
    ))
    stub_llm.reply("中的 Reviewer", json.dumps(
        {"passed": True, "reason": "ok", "missing_tasks": []},
        ensure_ascii=False,
    ))
    stub_llm.reply("Chart Planner", json.dumps(
        {"need_chart": False, "charts": [], "reason": "不需要"},
        ensure_ascii=False,
    ))
    stub_llm.reply("专业的数据分析师", "## 结论\n销售额合计 442,700 元。")

    # 每次分析消耗两条：一次工具调用 + 一次收尾
    for _ in range(4):

        tool_model.push_tool(fake_tool_call(
            "call_1",
            "aggregate_metric",
            {"group_by": None, "metric": "销售额", "agg": "sum"},
        ))

        tool_model.push_text("销售额合计 442,700 元。")

    return build_graph()


def test_the_graph_runs_the_full_chain_for_an_excel_upload(
    stubbed_graph, stub_llm, xlsx_file, make_state, trace_titles
):
    final = stubbed_graph.invoke(make_state(
        question="销售额是多少",
        file_path=xlsx_file,
        dataset_info=None,
    ))

    # Planner 跑过
    assert any("中的 Planner" in call for call in stub_llm.calls), stub_llm.calls

    assert final["input_kind"] == "table"
    assert final["source_type"] == "excel"

    # Executor 跑过，而且工具读的是 Excel 归一化出来的那份 CSV
    assert final["execution_results"]
    assert SALES_TOTAL in final["execution_results"][0]["data"]["result"]

    assert final["file_path"].endswith(".csv")
    assert os.path.isfile(final["file_path"])

    assert "Excel" in final["dataset_info"]
    assert any("Excel" in str(title) for title in trace_titles(final))

    # 报告节点也跑过
    assert any(
        "专业的数据分析师" in call for call in stub_llm.calls
    ), stub_llm.calls
    assert final["final_report"]


def test_the_graph_answers_word_prose_without_running_analysis(
    stubbed_graph, stub_llm, docx_prose_file, make_state
):
    final = stubbed_graph.invoke(make_state(
        question="会议讲了什么",
        file_path=docx_prose_file,
        dataset_info=None,
    ))

    assert final["input_kind"] == "text"
    assert final["source_type"] == "word_text"

    # 绕过 Planner / Executor
    assert not any("中的 Planner" in call for call in stub_llm.calls), stub_llm.calls
    assert not (final["execution_results"] or [])

    # 正文进了 state，也进了报告节点的提示词
    assert PROSE_LINE in final["input_context"]
    assert any(PROSE_LINE in call for call in stub_llm.calls), stub_llm.calls
    assert final["final_report"]


def test_the_graph_explains_an_unsupported_upload_without_analysis(
    stubbed_graph, stub_llm, pdf_file, make_state, trace_events
):
    final = stubbed_graph.invoke(make_state(
        question="销售额是多少",
        file_path=pdf_file,
        dataset_info=None,
    ))

    assert final["input_kind"] == "error"

    assert not any("中的 Planner" in call for call in stub_llm.calls), stub_llm.calls
    assert not (final["execution_results"] or [])

    # 报告如实说明读取失败，trace 里也留了记录
    assert "数据集读取失败" in final["final_report"]
    assert "dataset_error" in trace_events(final)


# =========================================================
# D. HTTP 契约：POST /analyze
# =========================================================

@pytest.fixture
def api_client(isolated_temp, stub_llm, tool_model, fake_tool_call):
    """
    ``/analyze`` 的测试客户端。

    模型全部换成桩；上传文件落在 ``isolated_temp`` 里，
    归一化 CSV 落在 ``tmp_path`` 里。
    """

    from fastapi.testclient import TestClient

    import main

    stub_llm.reply("中的 Planner", json.dumps(
        {"plan": [{"step": 1, "task": "汇总销售额", "tool": "run_stat_calc"}]},
        ensure_ascii=False,
    ))
    stub_llm.reply("中的 Reviewer", json.dumps(
        {"passed": True, "reason": "ok", "missing_tasks": []},
        ensure_ascii=False,
    ))
    stub_llm.reply("Chart Planner", json.dumps(
        {"need_chart": False, "charts": [], "reason": "不需要"},
        ensure_ascii=False,
    ))
    stub_llm.reply("专业的数据分析师", "## 结论\n销售额合计 442,700 元。")

    for _ in range(6):

        tool_model.push_tool(fake_tool_call(
            "call_1",
            "aggregate_metric",
            {"group_by": None, "metric": "销售额", "agg": "sum"},
        ))

        tool_model.push_text("销售额合计 442,700 元。")

    return TestClient(main.app)


def _post(client, name, path, mime):

    with open(path, "rb") as handle:

        return client.post(
            "/analyze",
            files={"file": (name, handle, mime)},
            data={"question": "各地区的销售额是多少"},
        )


def test_post_xlsx_returns_the_table_metadata(api_client, xlsx_file):
    response = _post(api_client, "sales.xlsx", xlsx_file, MIME_XLSX)

    assert response.status_code == 200

    payload = response.json()

    assert not payload.get("error"), str(payload.get("error"))[:200]
    assert payload["input_kind"] == "table"
    assert payload["source_type"] == "excel"
    assert isinstance(payload["dataset_notes"], list), payload.get("dataset_notes")


def test_post_docx_table_reports_a_word_table(api_client, docx_table_file):
    response = _post(
        api_client, "report_with_table.docx", docx_table_file, MIME_DOCX
    )

    assert response.status_code == 200
    assert response.json()["source_type"] == "word_table"


def test_post_an_unsupported_file_returns_a_friendly_error(api_client, pdf_file):
    response = _post(api_client, "notes.pdf", pdf_file, "application/pdf")

    payload = response.json()

    assert payload.get("error"), payload
    assert "支持的格式" in payload["error"]
    assert supported_hint() in payload["error"]


def test_post_a_plain_csv_still_works(api_client, csv_file):
    response = _post(api_client, "sales.csv", csv_file, "text/csv")

    assert response.status_code == 200

    payload = response.json()

    assert not payload.get("error"), str(payload.get("error"))[:200]
    assert payload["source_type"] == "csv"


def test_uploads_keep_their_original_extension(
    api_client, isolated_temp, xlsx_file, docx_table_file
):
    _post(api_client, "sales.xlsx", xlsx_file, MIME_XLSX)
    _post(api_client, "report_with_table.docx", docx_table_file, MIME_DOCX)

    saved = os.listdir(isolated_temp["temp"])

    assert any(
        name.startswith("data_") and name.endswith(".xlsx") for name in saved
    ), saved

    assert any(
        name.startswith("data_") and name.endswith(".docx") for name in saved
    ), saved
